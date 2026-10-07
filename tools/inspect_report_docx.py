from __future__ import annotations

import json
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


def length_text(text: str) -> int:
    return len("".join(text.split()))


def main() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    path = Path(sys.argv[1])
    document = Document(path)
    payload: dict[str, object] = {
        "path": str(path),
        "paragraph_count": len(document.paragraphs),
        "table_count": len(document.tables),
        "inline_shape_count": len(document.inline_shapes),
        "sections": [],
        "styles": {},
        "paragraphs": [],
        "tables": [],
        "headers_footers": [],
        "core_properties": {},
    }

    for section in document.sections:
        payload["sections"].append(
            {
                "page_width_cm": round(section.page_width.cm, 2),
                "page_height_cm": round(section.page_height.cm, 2),
                "top_cm": round(section.top_margin.cm, 2),
                "bottom_cm": round(section.bottom_margin.cm, 2),
                "left_cm": round(section.left_margin.cm, 2),
                "right_cm": round(section.right_margin.cm, 2),
                "header_cm": round(section.header_distance.cm, 2),
                "footer_cm": round(section.footer_distance.cm, 2),
            }
        )

    for style_name in ["Normal", "Title", "Heading 1", "Heading 2", "Caption"]:
        if style_name not in document.styles:
            continue
        style = document.styles[style_name]
        font = style.font
        rpr = style.element.rPr
        east_asia = None
        if rpr is not None and rpr.rFonts is not None:
            east_asia = rpr.rFonts.get(qn("w:eastAsia"))
        payload["styles"][style_name] = {
            "font": font.name,
            "east_asia": east_asia,
            "size_pt": font.size.pt if font.size else None,
            "bold": font.bold,
            "line_spacing": (
                style.paragraph_format.line_spacing.pt
                if getattr(style.paragraph_format.line_spacing, "pt", None)
                else style.paragraph_format.line_spacing
            ),
        }

    for index, paragraph in enumerate(document.paragraphs):
        text = paragraph.text.strip()
        if not text:
            continue
        payload["paragraphs"].append(
            {
                "index": index,
                "style": paragraph.style.name if paragraph.style else None,
                "length": length_text(text),
                "text": text,
            }
        )

    for table_index, table in enumerate(document.tables):
        rows = []
        for row in table.rows:
            rows.append([cell.text.strip() for cell in row.cells])
        payload["tables"].append({"index": table_index, "rows": rows})

    for section_index, section in enumerate(document.sections):
        payload["headers_footers"].append(
            {
                "section": section_index,
                "header": [p.text for p in section.header.paragraphs],
                "footer": [p.text for p in section.footer.paragraphs],
            }
        )

    props = document.core_properties
    payload["core_properties"] = {
        "title": props.title,
        "subject": props.subject,
        "author": props.author,
        "last_modified_by": props.last_modified_by,
        "keywords": props.keywords,
        "comments": props.comments,
    }

    print(json.dumps(payload, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
